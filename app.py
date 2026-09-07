from flask import Flask, render_template, request, jsonify, session, redirect, url_for, Response, send_file
import os
import json
from dotenv import load_dotenv
import threading
import time
from datetime import datetime
from db import db_get, db_insert, db_update, db_delete
import secrets
from werkzeug.middleware.proxy_fix import ProxyFix

load_dotenv()

app = Flask(__name__)

app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1, x_prefix=1)

os.environ['OAUTHLIB_INSECURE_TRANSPORT'] = '1'

app.secret_key = os.getenv("SECRET_KEY", "formgen-static-production-secret-key-9988")
app.config['PERMANENT_SESSION_LIFETIME'] = 60 * 60 * 24 * 30
app.config['SESSION_PERMANENT'] = True
app.config['SESSION_TYPE'] = 'filesystem'
app.config['SESSION_FILE_DIR'] = os.path.join(os.path.dirname(__file__), 'flask_sessions')
app.config['SESSION_FILE_THRESHOLD'] = 500
app.config['SESSION_COOKIE_SAMESITE'] = 'Lax'
app.config['SESSION_COOKIE_NAME'] = 'formgen_session'

#app.config['SESSION_COOKIE_DOMAIN'] = '127.0.0.1'

from flask_session import Session
Session(app)

# ── Google OAuth config ──────────────────────────────────────────
GOOGLE_CLIENT_ID = os.getenv("GOOGLE_CLIENT_ID")
GOOGLE_CLIENT_SECRET = os.getenv("GOOGLE_CLIENT_SECRET")
GOOGLE_REDIRECT_URI = os.getenv("GOOGLE_REDIRECT_URI", "http://127.0.0.1:5001/auth/callback")

# ── Helpers ──────────────────────────────────────────────────────
def get_current_user():
    return session.get('user')

def login_required(f):
    from functools import wraps
    @wraps(f)
    def decorated(*args, **kwargs):
        if not get_current_user():
            return redirect('/login')
        return f(*args, **kwargs)
    return decorated

def get_user_creds():
    from google.oauth2.credentials import Credentials
    token_data = session.get('google_token')
    if not token_data:
        return None
    return Credentials(
        token=token_data['access_token'],
        refresh_token=token_data.get('refresh_token'),
        token_uri='https://oauth2.googleapis.com/token',
        client_id=GOOGLE_CLIENT_ID,
        client_secret=GOOGLE_CLIENT_SECRET,
        scopes=token_data.get('scopes', [])
    )

def get_gmail_service():
    from googleapiclient.discovery import build
    creds = get_user_creds()
    if not creds:
        return None
    return build('gmail', 'v1', credentials=creds)

def send_notification_email(to_email, form_title, response_count):
    import base64
    from email.mime.text import MIMEText
    try:
        gmail_service = get_gmail_service()
        if not gmail_service:
            return False
        message_text = f"""
Hi,

Your form "{form_title}" has reached {response_count} responses.

You can view your form and responses here:
https://docs.google.com/forms/d/

This is an automated notification from FormGen.
"""
        message = MIMEText(message_text)
        message['to'] = to_email
        message['subject'] = f'FormGen: "{form_title}" has reached {response_count} responses'
        raw = base64.urlsafe_b64encode(message.as_bytes()).decode()
        gmail_service.users().messages().send(userId='me', body={'raw': raw}).execute()
        print(f"Notification email sent to {to_email}")
        return True
    except Exception as e:
        print(f"Error sending email: {e}")
        return False

from google_auth_oauthlib.flow import Flow

def get_google_oauth_flow():
    # Replace this dictionary with your existing Flow.from_client_config or client_secrets initialization
    return Flow.from_client_config(
        {
            "web": {
                "client_id": os.getenv("GOOGLE_CLIENT_ID"),
                "client_secret": os.getenv("GOOGLE_CLIENT_SECRET"),
                "auth_uri": "https://accounts.google.com/o/oauth2/auth",
                "token_uri": "https://oauth2.googleapis.com/token",
            }
        },
        scopes=[
            "openid",
            "https://www.googleapis.com/auth/userinfo.email",
            "https://www.googleapis.com/auth/userinfo.profile",
            "https://www.googleapis.com/auth/forms.body",
            "https://www.googleapis.com/auth/gmail.send"
        ]
    )

# ── Auth routes ───────────────────────────────────────────────────
@app.route('/login')
def login():
    error = request.args.get('error')
    if error:
        return render_template('login.html', error=error)

    redirect_uri = os.getenv("GOOGLE_REDIRECT_URI", "https://formgen-9pjf.onrender.com/auth/callback")
    flow = get_google_oauth_flow()
    flow.redirect_uri = redirect_uri

    authorization_url, state = flow.authorization_url(
        access_type='offline',
        prompt='select_account',
        include_granted_scopes='true'
    )
    
    session['oauth_state'] = state
    session['code_verifier'] = flow.code_verifier

    return redirect(authorization_url)

@app.route('/auth/google')
def auth_google():
    from google_auth_oauthlib.flow import Flow
    flow = Flow.from_client_config(
        {
            "web": {
                "client_id": GOOGLE_CLIENT_ID,
                "client_secret": GOOGLE_CLIENT_SECRET,
                "auth_uri": "https://accounts.google.com/o/oauth2/auth",
                "token_uri": "https://oauth2.googleapis.com/token",
                "redirect_uris": [GOOGLE_REDIRECT_URI]
            }
        },
        scopes=[
            'openid',
            'https://www.googleapis.com/auth/userinfo.email',
            'https://www.googleapis.com/auth/userinfo.profile',
            'https://www.googleapis.com/auth/forms.body',
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/gmail.send',
            'https://www.googleapis.com/auth/gmail.readonly'
        ]
    )
    flow.redirect_uri = GOOGLE_REDIRECT_URI
    auth_url, state = flow.authorization_url(
        access_type='offline',
        include_granted_scopes='true',
        prompt='consent'
    )
    session['oauth_state'] = state
    return redirect(auth_url)

@app.route('/auth/callback')
def auth_callback():
    from google_auth_oauthlib.flow import Flow
    import requests as req

    state = session.get('oauth_state')
    code_verifier = session.get('code_verifier')

    flow = Flow.from_client_config(
        {
            "web": {
                "client_id": GOOGLE_CLIENT_ID,
                "client_secret": GOOGLE_CLIENT_SECRET,
                "auth_uri": "https://accounts.google.com/o/oauth2/auth",
                "token_uri": "https://oauth2.googleapis.com/token",
                "redirect_uris": [GOOGLE_REDIRECT_URI]
            }
        },
        scopes=[
            'openid',
            'https://www.googleapis.com/auth/userinfo.email',
            'https://www.googleapis.com/auth/userinfo.profile',
            'https://www.googleapis.com/auth/forms.body',
            'https://www.googleapis.com/auth/drive',
            'https://www.googleapis.com/auth/gmail.send',
            'https://www.googleapis.com/auth/gmail.readonly'
        ],
        state=state
    )
    flow.redirect_uri = GOOGLE_REDIRECT_URI

    try:
        auth_response = request.url.replace('http://', 'https://', 1)
        flow.fetch_token(
            authorization_response=auth_response,
            code_verifier=code_verifier
        )
    except Exception as e:
        print(f"OAuth Callback Error: {e}")
        return redirect('/login?error=Authentication+failed')

    credentials = flow.credentials

    user_info_resp = req.get(
        'https://www.googleapis.com/oauth2/v2/userinfo',
        headers={'Authorization': f'Bearer {credentials.token}'}
    )
    user_info = user_info_resp.json()

    email = user_info.get('email')
    name = user_info.get('name')
    picture = user_info.get('picture')

    existing = db_get('users', {'email': f'eq.{email}'})
    if not existing or not isinstance(existing, list) or len(existing) == 0:
        db_insert('users', {'email': email, 'name': name, 'picture': picture})

    session.permanent = True
    session['user'] = {'email': email, 'name': name, 'picture': picture}
    session['google_token'] = {
        'access_token': credentials.token,
        'refresh_token': credentials.refresh_token,
        'scopes': list(credentials.scopes) if credentials.scopes else []
    }

    return redirect('/')

@app.route('/logout')
def logout():
    session.clear()
    return redirect('/login')

# ── Main routes ───────────────────────────────────────────────────
@app.route('/')
@login_required
def index():
    user = get_current_user()
    forms = db_get('forms', {'created_by': f'eq.{user["email"]}', 'order': 'created_on.desc'})
    if not isinstance(forms, list):
        forms = []

    memberships = db_get('workspace_members', {'email': f'eq.{user["email"]}', 'status': 'eq.approved'})
    user_workspaces = []
    if isinstance(memberships, list):
        for m in memberships:
            ws = db_get('workspaces', {'workspace_id': f'eq.{m["workspace_id"]}'})
            if ws and isinstance(ws, list):
                user_workspaces.append(ws[0])

    last_synced = datetime.now().strftime('%d %b %Y, %I:%M %p')
    return render_template('index.html', forms=forms, last_synced=last_synced, user=user, user_workspaces=user_workspaces)

@app.route('/create', methods=['POST'])
@login_required
def create_form():
    from main import parse_questions_with_ai, build_requests, enable_quiz_and_set_answers
    from googleapiclient.discovery import build

    user = get_current_user()
    title = request.form.get('title')
    user_input = request.form.get('questions')
    mandatory_choice = request.form.get('mandatory')
    expiry_time = request.form.get('expiry')
    threshold = request.form.get('threshold')
    allow_multiple = request.form.get('allow_multiple', 'no')

    if mandatory_choice == "1":
        mandatory_instruction = "ALL questions must be marked as required/mandatory."
    elif mandatory_choice == "2":
        mandatory_instruction = "NO questions should be marked as required/mandatory."
    else:
        mandatory_instruction = "Check the input carefully — if a question is marked as mandatory or required, set it as required. Otherwise leave it optional."

    creds = get_user_creds()
    service = build('forms', 'v1', credentials=creds)

    parsed = parse_questions_with_ai(user_input, title, mandatory_instruction)

    form = service.forms().create(body={"info": {"title": parsed["title"]}}).execute()
    form_id = form['formId']

    requests_list = build_requests(parsed["questions"])
    service.forms().batchUpdate(formId=form_id, body={"requests": requests_list}).execute()

    is_quiz = any(q.get("correct_answer") for q in parsed["questions"])
    if is_quiz:
        enable_quiz_and_set_answers(service, form_id, parsed["questions"])

    if allow_multiple == 'no':
        try:
            service.forms().batchUpdate(
                formId=form_id,
                body={"requests": [{"updateSettings": {"settings": {"limitOneResponsePerUser": True}, "updateMask": "limitOneResponsePerUser"}}]}
            ).execute()
        except Exception as e:
            print(f"Error setting response limit: {e}")

    form_link = f"https://docs.google.com/forms/d/{form_id}/edit"
    workspace_id = request.form.get('workspace_id') or None

    db_insert('forms', {
        'form_id': form_id,
        'form_title': parsed["title"],
        'created_by': user['email'],
        'form_link': form_link,
        'status': 'active',
        'expires_at': expiry_time if expiry_time else None,
        'response_threshold': int(threshold) if threshold else None,
        'allow_multiple': allow_multiple,
        'response_count': 0,
        'workspace_id': workspace_id,
        'visibility': 'workspace' if workspace_id else 'personal'
    })

    return jsonify({"success": True, "form_id": form_id, "link": form_link, "title": parsed["title"]})

@app.route('/form/<form_id>')
@login_required
def form_detail(form_id):
    user = get_current_user()

    # Fetch form without creator restriction
    forms = db_get('forms', {'form_id': f'eq.{form_id}'})
    if not forms or not isinstance(forms, list) or len(forms) == 0:
        return "Form not found", 404
    form = forms[0]

    # Check access — creator or approved workspace member
    is_creator = form.get('created_by') == user['email']
    is_workspace_member = False
    if form.get('workspace_id'):
        membership = db_get('workspace_members', {
            'workspace_id': f'eq.{form["workspace_id"]}',
            'email': f'eq.{user["email"]}',
            'status': 'eq.approved'
        })
        is_workspace_member = bool(membership and isinstance(membership, list) and len(membership) > 0)

    if not is_creator and not is_workspace_member:
        return "Access denied", 403

    response_count = 0
    analytics = []

    try:
        from googleapiclient.discovery import build
        creds = get_user_creds()
        service = build('forms', 'v1', credentials=creds)

        form_data = service.forms().get(formId=form_id).execute()
        items = form_data.get('items', [])
        responses_data = service.forms().responses().list(formId=form_id).execute()
        all_responses = responses_data.get('responses', [])
        response_count = len(all_responses)

        for item in items:
            question_item = item.get('questionItem', {})
            question = question_item.get('question', {})
            question_id = question.get('questionId', '')
            question_title = item.get('title', '')

            if 'choiceQuestion' in question:
                q_type = question['choiceQuestion'].get('type', 'RADIO')
                display_type = 'MULTIPLE_CHOICE' if q_type == 'RADIO' else ('CHECKBOX' if q_type == 'CHECKBOX' else 'DROPDOWN')
            elif 'scaleQuestion' in question:
                display_type = 'LINEAR_SCALE'
            elif 'ratingQuestion' in question:
                display_type = 'RATING'
            elif 'dateQuestion' in question:
                display_type = 'DATE'
            elif 'timeQuestion' in question:
                display_type = 'TIME'
            elif 'textQuestion' in question:
                display_type = 'PARAGRAPH' if question['textQuestion'].get('paragraph') else 'SHORT_ANSWER'
            else:
                display_type = 'UNKNOWN'

            answers = []
            for response in all_responses:
                answers_map = response.get('answers', {})
                if question_id in answers_map:
                    for a in answers_map[question_id].get('textAnswers', {}).get('answers', []):
                        answers.append(a.get('value', ''))

            chart_data = None
            if display_type in ['MULTIPLE_CHOICE', 'CHECKBOX', 'DROPDOWN']:
                options = question.get('choiceQuestion', {}).get('options', [])
                option_labels = [opt.get('value', '') for opt in options]
                option_counts = {label: 0 for label in option_labels}
                for answer in answers:
                    if answer in option_counts:
                        option_counts[answer] += 1
                chart_data = {'labels': list(option_counts.keys()), 'counts': list(option_counts.values())}
            elif display_type in ['LINEAR_SCALE', 'RATING']:
                scale_q = question.get('scaleQuestion') or question.get('ratingQuestion', {})
                low = scale_q.get('low', 1)
                high = scale_q.get('high', scale_q.get('ratingScaleLevel', 5))
                scale_counts = {str(i): 0 for i in range(low, high + 1)}
                for answer in answers:
                    if answer in scale_counts:
                        scale_counts[answer] += 1
                avg = round(sum(int(a) for a in answers if a.isdigit()) / len(answers), 2) if answers else 0
                chart_data = {'labels': list(scale_counts.keys()), 'counts': list(scale_counts.values()), 'average': avg}

            analytics.append({
                'question_id': question_id,
                'title': question_title,
                'type': display_type,
                'answers': answers,
                'chart_data': chart_data,
                'answer_count': len(answers)
            })

    except Exception as e:
        print(f"Error fetching analytics: {e}")

    return render_template('result.html', form=form, response_count=response_count, analytics=analytics, user=user)

@app.route('/form/<form_id>/export/csv')
@login_required
def export_csv(form_id):
    import csv
    import io

    try:
        from googleapiclient.discovery import build
        creds = get_user_creds()
        service = build('forms', 'v1', credentials=creds)

        form_data = service.forms().get(formId=form_id).execute()
        items = form_data.get('items', [])
        all_responses = service.forms().responses().list(formId=form_id).execute().get('responses', [])

        headers = ['Response #', 'Submitted At']
        question_ids = []
        for item in items:
            question = item.get('questionItem', {}).get('question', {})
            qid = question.get('questionId', '')
            if qid:
                headers.append(item.get('title', ''))
                question_ids.append(qid)

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)

        for i, response in enumerate(all_responses):
            row = [i + 1, response.get('lastSubmittedTime', '')]
            answers_map = response.get('answers', {})
            for qid in question_ids:
                if qid in answers_map:
                    row.append(', '.join([a.get('value', '') for a in answers_map[qid].get('textAnswers', {}).get('answers', [])]))
                else:
                    row.append('')
            writer.writerow(row)

        output.seek(0)
        return Response(
            output.getvalue(),
            mimetype='text/csv',
            headers={'Content-Disposition': f'attachment; filename=Responses_{form_data.get("info", {}).get("title", "form").replace(" ", "_")}.csv'}
        )
    except Exception as e:
        return f"Error: {e}", 500

@app.route('/form/<form_id>/export/excel')
@login_required
def export_excel(form_id):
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    try:
        from googleapiclient.discovery import build
        creds = get_user_creds()
        service = build('forms', 'v1', credentials=creds)

        form_data = service.forms().get(formId=form_id).execute()
        items = form_data.get('items', [])
        all_responses = service.forms().responses().list(formId=form_id).execute().get('responses', [])

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Responses"

        header_fill = PatternFill(start_color="1a1a1a", end_color="1a1a1a", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True)

        headers = ['Response #', 'Submitted At']
        question_ids = []
        for item in items:
            question = item.get('questionItem', {}).get('question', {})
            qid = question.get('questionId', '')
            if qid:
                headers.append(item.get('title', ''))
                question_ids.append(qid)

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center')

        for i, response in enumerate(all_responses):
            row_data = [i + 1, response.get('lastSubmittedTime', '')]
            answers_map = response.get('answers', {})
            for qid in question_ids:
                if qid in answers_map:
                    row_data.append(', '.join([a.get('value', '') for a in answers_map[qid].get('textAnswers', {}).get('answers', [])]))
                else:
                    row_data.append('')
            for col, value in enumerate(row_data, 1):
                ws.cell(row=i + 2, column=col, value=value)

        for col in ws.columns:
            max_length = max((len(str(cell.value or '')) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 4, 50)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        return send_file(
            output,
            mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            as_attachment=True,
            download_name=f'Responses_{form_data.get("info", {}).get("title", "form").replace(" ", "_")}.xlsx'
        )
    except Exception as e:
        return f"Error: {e}", 500

@app.route('/form/<form_id>/close', methods=['POST'])
@login_required
def close_form(form_id):
    try:
        from googleapiclient.discovery import build
        creds = get_user_creds()
        service = build('forms', 'v1', credentials=creds)
        service.forms().setPublishSettings(
            formId=form_id,
            body={"publishSettings": {"publishState": {"isPublished": True, "isAcceptingResponses": False}}}
        ).execute()
        db_update('forms', {'form_id': f'eq.{form_id}'}, {'status': 'closed'})
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})

@app.route('/form/<form_id>/reopen', methods=['POST'])
@login_required
def reopen_form(form_id):
    try:
        from googleapiclient.discovery import build
        creds = get_user_creds()
        service = build('forms', 'v1', credentials=creds)
        service.forms().setPublishSettings(
            formId=form_id,
            body={"publishSettings": {"publishState": {"isPublished": True, "isAcceptingResponses": True}}}
        ).execute()
        clear_expiry = request.json.get('clear_expiry', False)
        update_data = {'status': 'active'}
        if clear_expiry:
            update_data['expires_at'] = None
        db_update('forms', {'form_id': f'eq.{form_id}'}, update_data)
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})

@app.route('/form/<form_id>/delete', methods=['POST'])
@login_required
def delete_form(form_id):
    try:
        from googleapiclient.discovery import build
        creds = get_user_creds()
        drive_service = build('drive', 'v3', credentials=creds)
        drive_service.files().delete(fileId=form_id).execute()
        db_delete('forms', {'form_id': f'eq.{form_id}'})
        return jsonify({"success": True})
    except Exception as e:
        return jsonify({"success": False, "error": str(e)})

@app.route('/settings', methods=['GET'])
@login_required
def settings():
    user = get_current_user()
    settings_list = db_get('settings', {'email': f'eq.{user["email"]}'})
    settings_data = settings_list[0] if settings_list and isinstance(settings_list, list) and len(settings_list) > 0 else {}
    return render_template('settings.html', settings=settings_data, user=user)

@app.route('/settings/save', methods=['POST'])
@login_required
def save_settings():
    user = get_current_user()
    notification_email = request.form.get('notification_email', '')
    response_threshold = int(request.form.get('response_threshold', 10))

    existing = db_get('settings', {'email': f'eq.{user["email"]}'})
    if existing and isinstance(existing, list) and len(existing) > 0:
        db_update('settings', {'email': f'eq.{user["email"]}'}, {
            'notification_email': notification_email,
            'response_threshold': response_threshold
        })
    else:
        db_insert('settings', {
            'email': user['email'],
            'notification_email': notification_email,
            'response_threshold': response_threshold
        })
    return jsonify({"success": True})

# ── Workspace routes ──────────────────────────────────────────────

@app.route('/workspaces')
@login_required
def workspaces():
    user = get_current_user()
    email = user['email']

    owned = db_get('workspaces', {'created_by': f'eq.{email}'})
    if not isinstance(owned, list):
        owned = []

    memberships = db_get('workspace_members', {'email': f'eq.{email}', 'status': 'eq.approved'})
    if not isinstance(memberships, list):
        memberships = []

    owned_ids = [w['workspace_id'] for w in owned]
    joined = []
    for m in memberships:
        wid = m['workspace_id']
        if wid not in owned_ids:
            ws = db_get('workspaces', {'workspace_id': f'eq.{wid}'})
            if ws and isinstance(ws, list):
                joined.append(ws[0])

    return render_template('workspaces.html', owned=owned, joined=joined, user=user)


@app.route('/workspaces/create', methods=['POST'])
@login_required
def create_workspace():
    import uuid
    user = get_current_user()
    name = request.form.get('name', '').strip()
    if not name:
        return jsonify({"success": False, "error": "Workspace name is required"})

    existing = db_get('workspaces', {'created_by': f'eq.{user["email"]}', 'name': f'eq.{name}'})
    if existing and isinstance(existing, list) and len(existing) > 0:
        return jsonify({"success": False, "error": f'You already have a workspace named "{name}"'})

    invite_code = str(uuid.uuid4())[:8].upper()

    result = db_insert('workspaces', {
        'name': name,
        'invite_code': invite_code,
        'created_by': user['email']
    })

    print("Create workspace result:", result)

    if isinstance(result, list) and len(result) > 0:
        workspace_id = result[0].get('workspace_id') or result[0].get('id')
    elif isinstance(result, dict):
        workspace_id = result.get('workspace_id') or result.get('id')
    else:
        workspace_id = None

    if not workspace_id:
        return jsonify({"success": False, "error": f"Could not get workspace ID: {result}"})

    db_insert('workspace_members', {
        'workspace_id': workspace_id,
        'email': user['email'],
        'role': 'admin',
        'status': 'approved'
    })

    return jsonify({"success": True, "workspace_id": workspace_id, "invite_code": invite_code})


@app.route('/workspaces/join', methods=['POST'])
@login_required
def join_workspace():
    user = get_current_user()
    invite_code = request.form.get('invite_code', '').strip().upper()

    ws = db_get('workspaces', {'invite_code': f'eq.{invite_code}'})
    if not ws or not isinstance(ws, list) or len(ws) == 0:
        return jsonify({"success": False, "error": "Invalid invite code"})

    workspace = ws[0]
    workspace_id = workspace['workspace_id']

    existing = db_get('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{user["email"]}'
    })
    if existing and isinstance(existing, list) and len(existing) > 0:
        status = existing[0]['status']
        if status == 'approved':
            return jsonify({"success": False, "error": "You are already a member"})
        elif status == 'pending':
            return jsonify({"success": False, "error": "Your request is already pending approval"})

    db_insert('workspace_members', {
        'workspace_id': workspace_id,
        'email': user['email'],
        'role': 'member',
        'status': 'pending'
    })

    return jsonify({"success": True, "workspace_name": workspace['name']})


@app.route('/workspaces/<int:workspace_id>')
@login_required
def workspace_detail(workspace_id):
    user = get_current_user()
    email = user['email']

    ws = db_get('workspaces', {'workspace_id': f'eq.{workspace_id}'})
    if not ws or not isinstance(ws, list) or len(ws) == 0:
        return "Workspace not found", 404
    workspace = ws[0]

    membership = db_get('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{email}'
    })
    if not membership or not isinstance(membership, list) or len(membership) == 0:
        return "Access denied", 403
    member = membership[0]
    if member['status'] != 'approved':
        return "Your membership is pending approval", 403

    is_admin = member['role'] == 'admin'

    members = db_get('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'status': 'eq.approved'
    })
    if not isinstance(members, list):
        members = []

    pending = []
    if is_admin:
        pending = db_get('workspace_members', {
            'workspace_id': f'eq.{workspace_id}',
            'status': 'eq.pending'
        })
        if not isinstance(pending, list):
            pending = []

    forms = db_get('forms', {
        'workspace_id': f'eq.{workspace_id}',
        'order': 'created_on.desc'
    })
    if not isinstance(forms, list):
        forms = []

    return render_template('workspace_detail.html',
        workspace=workspace,
        members=members,
        pending=pending,
        forms=forms,
        is_admin=is_admin,
        user=user
    )


@app.route('/workspaces/<int:workspace_id>/approve', methods=['POST'])
@login_required
def approve_member(workspace_id):
    user = get_current_user()
    membership = db_get('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{user["email"]}',
        'role': 'eq.admin'
    })
    if not membership or not isinstance(membership, list) or len(membership) == 0:
        return jsonify({"success": False, "error": "Not authorized"})

    member_email = request.json.get('email')
    db_update('workspace_members',
        {'workspace_id': f'eq.{workspace_id}', 'email': f'eq.{member_email}'},
        {'status': 'approved'}
    )
    return jsonify({"success": True})


@app.route('/workspaces/<int:workspace_id>/reject', methods=['POST'])
@login_required
def reject_member(workspace_id):
    user = get_current_user()
    membership = db_get('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{user["email"]}',
        'role': 'eq.admin'
    })
    if not membership or not isinstance(membership, list) or len(membership) == 0:
        return jsonify({"success": False, "error": "Not authorized"})

    member_email = request.json.get('email')
    db_delete('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{member_email}'
    })
    return jsonify({"success": True})


@app.route('/workspaces/<int:workspace_id>/remove', methods=['POST'])
@login_required
def remove_member(workspace_id):
    user = get_current_user()
    membership = db_get('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{user["email"]}',
        'role': 'eq.admin'
    })
    if not membership or not isinstance(membership, list) or len(membership) == 0:
        return jsonify({"success": False, "error": "Not authorized"})

    member_email = request.json.get('email')
    db_delete('workspace_members', {
        'workspace_id': f'eq.{workspace_id}',
        'email': f'eq.{member_email}'
    })
    return jsonify({"success": True})


@app.route('/workspaces/<int:workspace_id>/delete', methods=['POST'])
@login_required
def delete_workspace(workspace_id):
    user = get_current_user()
    ws = db_get('workspaces', {'workspace_id': f'eq.{workspace_id}', 'created_by': f'eq.{user["email"]}'})
    if not ws or not isinstance(ws, list) or len(ws) == 0:
        return jsonify({"success": False, "error": "Not authorized"})
    db_delete('workspaces', {'workspace_id': f'eq.{workspace_id}'})
    return jsonify({"success": True})


# ── Startup ───────────────────────────────────────────────────────
if __name__ == '__main__':
    app.run(debug=True, port=5001, use_reloader=False, host='0.0.0.0')